"""Formula-based independent checks with evaluated Excel formula caches.

The restricted evaluator executes the saved spreadsheet formulas, not the
hydrology or routing functions. Excel recalculates them when inputs change.
It is deliberately limited to the arithmetic and functions this workbook uses.
"""
import ast
import json
import math
from pathlib import Path
import re

import numpy as np
import openpyxl
import pandas as pd
import xlsxwriter


class FormulaEvaluator:
    def __init__(self, cells):
        self.cells = cells
        self.cache = {}
        self.active = set()

    def cell(self, sheet, address):
        key = sheet, address.replace("$", "")
        if key in self.cache:
            return self.cache[key]
        value = self.cells.get(key, 0)
        if isinstance(value, str) and value.startswith("="):
            if key in self.active:
                raise ValueError(f"Circular formula: {key}")
            self.active.add(key)
            value = self.evaluate(value, sheet)
            self.active.remove(key)
        self.cache[key] = value
        return value

    def evaluate(self, formula, sheet):
        expression = formula[1:].replace("^", "**").replace("<>", "!=")
        expression = re.sub(r"(?<![<>=!])=(?!=)", "==", expression)
        references = []
        def replace(match):
            reference_sheet = match.group(1) or match.group(2) or sheet
            address = match.group(3).replace("$", "")
            references.append((reference_sheet, address))
            return f"REF({len(references)-1})"
        expression = re.sub(r"(?:(?:'([^']+)'|([A-Za-z_][A-Za-z_0-9]*))!)?(\$?[A-Z]{1,3}\$?[0-9]+)\b", replace, expression)
        node = ast.parse(expression, mode="eval").body
        def visit(n):
            if isinstance(n, ast.Constant) and isinstance(n.value, (int, float, bool, str)):
                return n.value
            if isinstance(n, ast.UnaryOp):
                v = visit(n.operand)
                if isinstance(n.op, ast.USub): return -v
                if isinstance(n.op, ast.UAdd): return v
            if isinstance(n, ast.BinOp):
                a, b = visit(n.left), visit(n.right)
                if isinstance(n.op, ast.Add): return a + b
                if isinstance(n.op, ast.Sub): return a - b
                if isinstance(n.op, ast.Mult): return a * b
                if isinstance(n.op, ast.Div): return a / b
                if isinstance(n.op, ast.Pow): return a**b
            if isinstance(n, ast.Compare) and len(n.ops) == 1:
                a, b = visit(n.left), visit(n.comparators[0])
                op = n.ops[0]
                if isinstance(op, ast.Lt): return a < b
                if isinstance(op, ast.LtE): return a <= b
                if isinstance(op, ast.Gt): return a > b
                if isinstance(op, ast.GtE): return a >= b
                if isinstance(op, ast.Eq): return a == b
                if isinstance(op, ast.NotEq): return a != b
            if isinstance(n, ast.Call) and isinstance(n.func, ast.Name):
                name = n.func.id.upper()
                if name == "REF":
                    return self.cell(*references[visit(n.args[0])])
                if name == "IF":
                    return visit(n.args[1] if visit(n.args[0]) else n.args[2])
                args = [visit(a) for a in n.args]
                functions = {"MAX": max, "MIN": min, "ABS": lambda a: abs(a[0]),
                             "SQRT": lambda a: math.sqrt(a[0]), "SUM": sum,
                             "AND": all, "OR": any}
                if name in functions:
                    return functions[name](args)
            raise ValueError(f"Unsupported spreadsheet formula: {formula}, {ast.dump(n)}")
        return visit(node)


def build_workbook(root, config, cases, baselines, selection, sensitivity):
    root = Path(root)
    wb = xlsxwriter.Workbook(root / "stormwater_calculations.xlsx")
    wb.set_properties({"title": "Stormwater detention calculations", "author": "Codex-assisted conceptual study",
                       "comments": "Synthetic event study. Independent formula checks are separate from imported scenario results."})
    wb.set_calc_mode("auto")
    title = wb.add_format({"font_size": 20, "bold": True, "font_color": "#163640"})
    head = wb.add_format({"bold": True, "bg_color": "#163640", "font_color": "white", "text_wrap": True, "valign": "vcenter"})
    numeric = wb.add_format({"num_format": "0.000000", "font_color": "#174f45"})
    formula_fmt = wb.add_format({"num_format": "0.000000", "bg_color": "#e8f5ed"})
    note = wb.add_format({"text_wrap": True, "font_color": "#526470", "valign": "top"})
    cells, sheet_map = {}, {}
    def sheet(name, heading, description):
        ws = wb.add_worksheet(name)
        sheet_map[name] = ws
        ws.hide_gridlines(2)
        ws.set_landscape()
        ws.fit_to_pages(1, 0)
        ws.freeze_panes(4, 1)
        ws.set_column(0, 12, 19)
        ws.merge_range("A1:J1", heading, title)
        ws.merge_range("A2:J2", description, note)
        ws.set_row(1, 32)
        return ws
    def put(ws, address, value, fmt=None):
        cells[(ws.name, address)] = value
        if isinstance(value, str) and value.startswith("="):
            evaluator = FormulaEvaluator(cells)
            cached = evaluator.cell(ws.name, address)
            ws.write_formula(address, value, fmt or formula_fmt, cached)
        else:
            ws.write(address, value, fmt)
    def headings(ws, row, names):
        for col, name in enumerate(names):
            ws.write(row - 1, col, name, head)
        ws.set_row(row - 1, 32)
    winner = selection["selected"]
    example = winner or {"storage_depth_m": config["storage_depths_m"][-1],
                         "outlet_side_mm": config["outlet_sides_mm"][-1]}
    ws = sheet("Inputs", "Stormwater detention | calculation workbook",
        "Green cells contain independent formulas. Scenario tables are saved Python results; changing this workbook does not rerun the grid.")
    parameters = [
        (3, "Parking area", config["parking_area_m2"], "m²"),
        (4, "Basin footprint", config["basin_area_m2"], "m²"),
        (5, "Curve number", config["curve_number"], "assumed"),
        (6, "Initial abstraction / S", config["abstraction_ratio"], "ratio"),
        (7, "Lag", config["lag_min"], "min"),
        (8, "Discharge coefficient", config["discharge_coefficient"], "assumed"),
        (9, "Gravity", config["gravity_m_s2"], "m/s²"),
        (10, "Selected depth", example["storage_depth_m"], "m; example if infeasible"),
        (11, "Outlet side", example["outlet_side_mm"] / 1000, "m"),
        (12, "Final simulation step", selection["final_dt_s"], "s"),
        (13, "Storage capacity", "=B4*B10", "m³"),
        (14, "Design rainfall", config["design_rainfall_mm"], "mm"),
        (15, "Absolute check tolerance", 1e-8, "m³"),
    ]
    for row, label, value, unit in parameters:
        put(ws, f"A{row}", label)
        put(ws, f"B{row}", value, numeric if not isinstance(value, str) else None)
        put(ws, f"C{row}", unit)
    ws.set_column("A:A", 29)
    ws.set_column("C:C", 31)
    ws.merge_range("A18:J19", "Site area is 4,200 m² in every default case. Direct footprint rainfall joins the baseline discharge or basin storage. Zero basin infiltration; detention changes timing, not total event runoff.", note)
    ws = sheet("Runoff checks", "Independent CN volume calculation",
        "These cells calculate retention, initial abstraction, excess depth, and runoff volume directly from Excel formulas. Column H is only the comparison reference.")
    headings(ws, 4, ["Rain (mm)", "S (mm)", "Ia (mm)", "Excess (mm)", "Parking (m³)", "Footprint (m³)", "Total (m³)", "Python total (m³)", "Difference (m³)", "Pass"])
    for row, rain in enumerate(config["rainfall_depths_mm"], start=5):
        reference = float(baselines[np.isclose(baselines.rainfall_mm, rain)].total_inflow_m3.iloc[0])
        values = {"A": rain, "B": "=25400/Inputs!B5-254", "C": f"=Inputs!B6*B{row}",
                  "D": f"=IF(A{row}<=C{row},0,(A{row}-C{row})^2/(A{row}-C{row}+B{row}))",
                  "E": f"=D{row}*Inputs!B3/1000", "F": f"=A{row}*Inputs!B4/1000",
                  "G": f"=E{row}+F{row}", "H": reference, "I": f"=G{row}-H{row}",
                  "J": f"=ABS(I{row})<=Inputs!B15"}
        for col, value in values.items(): put(ws, f"{col}{row}", value)
    ws.merge_range("A11:J12", "Try CN 95 or 100 on Inputs. Columns B–G recompute; the Python reference stays fixed until run_project.py is rerun. Initial abstraction is applied once to cumulative rainfall. Depth × area / 1,000 converts mm to m³.", note)
    chart = wb.add_chart({"type": "column", "subtype": "stacked"})
    for col, name, color in [(4, "Parking runoff", "#157f79"), (5, "Direct footprint rainfall", "#86cab5")]:
        chart.add_series({"name": name, "categories": [ws.name, 4, 0, 3 + len(config["rainfall_depths_mm"]), 0],
            "values": [ws.name, 4, col, 3 + len(config["rainfall_depths_mm"]), col], "fill": {"color": color}})
    chart.set_title({"name": "Same volume enters every design"})
    chart.set_x_axis({"name": "Event rainfall (mm)"})
    chart.set_y_axis({"name": "Runoff + direct rainfall (m³)"})
    ws.insert_chart("A15", chart, {"x_scale": 1.4})
    ws = sheet("Storage check", "Independent continuity | blocked outlet",
        "A separate exact fixture: 1 m³/s for eight 5-second steps into a 20 m³ tank, with no controlled outlet. After filling, every extra m³ becomes overflow.")
    for address, value in {"A3": "Capacity (m³)", "B3": 20, "D3": "Step (s)", "E3": 5}.items(): put(ws, address, value)
    headings(ws, 5, ["Step", "Time end (s)", "Inflow (m³/s)", "Start S (m³)", "Available (m³)", "End S (m³)", "Overflow (m³)", "Cumulative overflow", "Step residual (m³)"])
    for row in range(6, 14):
        values = {"A": row - 5, "B": f"=A{row}*$E$3", "C": 1,
            "D": 0 if row == 6 else f"=F{row-1}", "E": f"=D{row}+C{row}*$E$3",
            "F": f"=MIN($B$3,E{row})", "G": f"=MAX(E{row}-$B$3,0)",
            "H": f"=G{row}" if row == 6 else f"=H{row-1}+G{row}",
            "I": f"=D{row}+C{row}*$E$3-F{row}-G{row}"}
        for col, value in values.items(): put(ws, f"{col}{row}", value)
    put(ws, "A16", "Total input (m³)"); put(ws, "B16", "=8*$E$3")
    put(ws, "A17", "Storage + overflow (m³)"); put(ws, "B17", "=F13+H13")
    put(ws, "A18", "Water balance pass"); put(ws, "B18", "=ABS(B16-B17)<=Inputs!B15")
    chart = wb.add_chart({"type": "line"})
    for col, label in ((5, "Storage"), (7, "Cumulative overflow")):
        chart.add_series({"name": label, "categories": [ws.name, 5, 1, 12, 1], "values": [ws.name, 5, col, 12, col]})
    chart.set_title({"name": "Clipped storage keeps overflow in the balance"})
    chart.set_x_axis({"name": "Time (s)"}); chart.set_y_axis({"name": "Volume (m³)"})
    ws.insert_chart("A21", chart, {"x_scale": 1.4})
    ws = sheet("Implicit step", "Independent Excel bisection | one storage step",
        "45 formula iterations solve Snew + dt × Q(Snew/A) = Sold + Vin. This reproduces a controlled-outlet step without calling Python's root finder.")
    fixture = {"B4": 200, "B5": 120, "B6": 40, "B7": 1.5, "B8": 30,
               "B9": 0.08, "B10": 0.62, "B11": 9.80665}
    labels = {4: "Area (m²)", 5: "Capacity (m³)", 6: "Old storage (m³)", 7: "Inflow volume (m³)",
              8: "Step (s)", 9: "Opening side (m)", 10: "Cd", 11: "Gravity (m/s²)"}
    for address, value in fixture.items(): put(ws, address, value)
    for row, label in labels.items(): put(ws, f"A{row}", label)
    headings(ws, 14, ["Iteration", "Lower S", "Upper S", "Midpoint S", "Depth (m)", "Outlet (m³/s)", "Equation residual (m³)"])
    for row in range(15, 60):
        values = {"A": row - 14,
            "B": 0 if row == 15 else f"=IF(G{row-1}>0,B{row-1},D{row-1})",
            "C": "=MIN($B$5,$B$6+$B$7)" if row == 15 else f"=IF(G{row-1}>0,D{row-1},C{row-1})",
            "D": f"=(B{row}+C{row})/2", "E": f"=D{row}/$B$4",
            "F": f"=(2/3)*$B$10*$B$9*SQRT(2*$B$11)*(E{row}^1.5-MAX(E{row}-$B$9,0)^1.5)",
            "G": f"=D{row}+$B$8*F{row}-($B$6+$B$7)"}
        for col, value in values.items(): put(ws, f"{col}{row}", value)
    put(ws, "I4", "Formula final S (m³)"); put(ws, "J4", "=D59")
    from .routing import route
    reference = route([1.5], 30, 200, 120, 0.08, initial_storage_m3=40).storage_m3[-1]
    put(ws, "I5", "Python comparison S"); put(ws, "J5", float(reference))
    put(ws, "I6", "Difference (m³)"); put(ws, "J6", "=J4-J5")
    put(ws, "I7", "Independent check pass"); put(ws, "J7", "=ABS(J6)<=Inputs!B15")
    put(ws, "I8", "Equation residual (m³)"); put(ws, "J8", "=G59")
    ws.set_column("A:A", 24)
    ws.set_column("I:I", 26)
    def dataframe_sheet(name, title_text, description, frame):
        ws = sheet(name, title_text, description)
        columns = list(frame.columns)
        for j, col in enumerate(columns): ws.write(3, j, col, head)
        for i, values in enumerate(frame.itertuples(index=False, name=None), start=4):
            for j, value in enumerate(values):
                if pd.isna(value): continue
                if isinstance(value, (np.integer, np.floating)): value = value.item()
                ws.write(i, j, value)
        if len(frame):
            ws.autofilter(3, 0, len(frame) + 3, len(columns) - 1)
        return ws
    dataframe_sheet("Scenarios", "All candidate-event results", "Imported numerical results at the verified final resolution; flow is m³/s and downstream discharge includes overflow.", cases)
    sel_frame = pd.read_csv(root / "results/design_selection.csv")
    ws = dataframe_sheet("Design selection", "Compare every tested design", "All storm shapes at the design depth must pass. Lowest capacity wins; shortest worst drawdown breaks a capacity tie.", sel_frame)
    chart = wb.add_chart({"type": "column"})
    col = list(sel_frame.columns).index("worst_peak_reduction_pct")
    chart.add_series({"name": "Worst peak reduction across design storms", "categories": [ws.name, 4, 0, len(sel_frame) + 3, 0],
                      "values": [ws.name, 4, col, len(sel_frame) + 3, col], "fill": {"color": "#157f79"}})
    chart.set_title({"name": "Peak reduction alone does not ensure no overflow"})
    chart.set_y_axis({"name": "Peak reduction (%)"}); chart.set_x_axis({"name": "Depth and opening configuration", "num_font": {"rotation": -45}})
    chart.set_size({"width": 1100, "height": 430}); ws.insert_chart(len(sel_frame) + 7, 0, chart)
    dataframe_sheet("Sensitivity", "One-at-a-time sensitivity", "Only the named parameter changes. Matching baseline runoff is recomputed for CN and lag changes; nominal rows repeat intentionally in each parameter group.", sensitivity)
    dataframe_sheet("Convergence", "Time-step verification", "Relative changes compare coarse and fine outputs. Overflow onset is the start of the first overflowing interval; drawdown uses the final threshold crossing.", pd.read_csv(root / "results/convergence.csv"))
    sources = json.loads((root / "results/source_inputs/sources.json").read_text())
    dataframe_sheet("Sources", "Source register", "All numerical source data are bundled with the project. URLs document provenance and are not needed for recomputation.", pd.DataFrame(sources))
    wb.close()
    return audit_workbook(root / "stormwater_calculations.xlsx")


def audit_workbook(path):
    formulas = openpyxl.load_workbook(path, data_only=False)
    cached = openpyxl.load_workbook(path, data_only=True)
    cells = {(ws.title, cell.coordinate): cell.value for ws in formulas for row in ws for cell in row if cell.value is not None}
    evaluator = FormulaEvaluator(cells)
    checked, errors = 0, []
    for (sheet, address), value in cells.items():
        if isinstance(value, str) and value.startswith("="):
            actual = evaluator.cell(sheet, address)
            cache = cached[sheet][address].value
            checked += 1
            if cache is None or not math.isclose(float(actual), float(cache), rel_tol=1e-12, abs_tol=1e-10):
                errors.append(f"{sheet}!{address}: formula={actual}, cache={cache}")
    pass_cells = [("Storage check", "B18"), ("Implicit step", "J7")]
    pass_cells += [("Runoff checks", f"J{r}") for r in range(5, formulas["Runoff checks"].max_row + 1)
                   if formulas["Runoff checks"][f"J{r}"].data_type == "f"]
    for sheet, cell in pass_cells:
        if evaluator.cell(sheet, cell) is not True:
            errors.append(f"Independent check did not pass: {sheet}!{cell}")
    return {"passed": not errors, "formulas_evaluated": checked, "errors": errors,
            "method": "Restricted independent evaluator re-read every saved Excel formula and checked every cache; Excel GUI was not used.",
            "independent_pass_cells": [f"{s}!{c}" for s, c in pass_cells],
            "charts": sum(len(ws._charts) for ws in formulas)}
